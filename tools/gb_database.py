"""The owner's product database (an Excel workbook kept OFF the site): every Gioberti product, what is on the site, and everything taken off it.

  Removed Products   listings left off the site (gb_catalog.HIDDEN / HIDDEN_CATS): photo, style, why and when removed, stock at Gioberti, colours,
                     sizes, price, links to the Gioberti page and the local photo folder (~/Downloads/Gioberti)
  Removed Colours    colourways trimmed from listings that are still on the site (gb_catalog.CURATED), with their photo and stock
  Full Catalog       every Gioberti listing: on the site / combined into another page / removed, with its site category and page
  Summary            counts (formulas) and how to bring something back

Reads Gioberti's live products.json (the same data as tools/gb_fetch.py) and the site rules in tools/gb_catalog.py; photos come from
assets/img/gioberti (run gb_fetch.py first if a photo is missing). The date each item was first seen removed is kept in removed-log.json
next to the workbook. Run:  python tools/gb_database.py ["C:/path/to/Checkmatela Product Database.xlsx"]
Default location: "Product Database" folder next to the repository (not deployed, not on GitHub).
"""
import os, sys, io, json, re, datetime
from PIL import Image as PILImage
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import gb_fetch
import gb_catalog as gc

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DEFAULT = os.path.join(os.path.dirname(ROOT), "Product Database", "Checkmatela Product Database.xlsx")
ASSETS = os.path.join(ROOT, "assets", "img", "gioberti")
SITE = "https://www.checkmatela.com/"
FONT = "Arial"
HEAD_FILL = PatternFill("solid", start_color="1E4D3B")             # the site's pine green
BAND_FILL = PatternFill("solid", start_color="F4F1E8")
THIN = Border(bottom=Side(style="thin", color="D9D4C7"))


def thumb(key, size=72):
    """A small PNG of a Gioberti photo for the sheet (Excel cannot show WebP)."""
    path = os.path.join(ASSETS, f"{key}-t.webp")
    if not key or not os.path.exists(path):
        return None
    im = PILImage.open(path).convert("RGB")
    im.thumbnail((size, size))
    buf = io.BytesIO()
    im.save(buf, "PNG")
    buf.seek(0)
    xl = XLImage(buf)
    xl.width, xl.height = im.size
    return xl


def stock(it, colour=None):
    n = len(it["options"][1]["values"]) if len(it["options"]) > 1 else 1
    rows = [r for r in it["variants"] if colour is None or r[0] == colour]
    ok = sum(1 for r in rows if r[-1])
    return ok, len(rows) if colour is None else n


def site_rules():
    """handle -> (status, where on the site / why)"""
    import gen_gioberti as gg                                          # same filing rules the site is built with
    status = {}
    for cfg in gc.CURATED:
        for h in cfg["sources"]:
            status[h] = ("Combined", f"{cfg['name']}  ({SITE}gioberti-{cfg['handle']}.html)")
    return status, gg


def main():
    args = [a for i, a in enumerate(sys.argv[1:], 1) if a != "--sheets-json" and sys.argv[i - 1] != "--sheets-json"]
    out = args[0] if args else DEFAULT
    os.makedirs(os.path.dirname(out), exist_ok=True)
    log_path = os.path.join(os.path.dirname(out), "removed-log.json")
    log = json.load(open(log_path, encoding="utf-8")) if os.path.exists(log_path) else {}
    today = datetime.date.today().isoformat()

    items, _ = gb_fetch.normalise(gb_fetch.fetch_products())
    gb_fetch.archive_paths(items)
    keys = {}                                                          # photo URL -> web copy key, from the site's catalog
    for it in json.load(open(os.path.join(ROOT, "tools", "catalog", "gioberti.json"), encoding="utf-8")):
        for im in it["images"]:
            if im.get("url") and im.get("key"):
                keys[im["url"]] = im["key"]
    for it in items:
        for im in it["images"]:
            im["key"] = keys.get(im["src"].split("?")[0])
    combined, gg = site_rules()

    def cat_label(h, it):
        c = gc.category_of(h, it["ptype"], it["name"])
        if not c:
            return "(no category)"
        dept, group, label, _ = gc.CATS[c]
        return f"{'Men' if dept == 'men' else 'Boys'} - {label}"

    def colour_photo(it, colour):
        by_id = {}
        for im in it["images"]:
            if colour in (im.get("colours") or []) and not im["chart"]:
                return im["key"]
        return next((im["key"] for im in it["images"] if not im["chart"] and im.get("key")), None)

    def folder(it, colour=None):
        files = [f for im in it["images"] for f in im["files"] if colour is None or colour in (im.get("colours") or [])] or [f for im in it["images"] for f in im["files"]]
        return os.path.dirname(files[0]) if files else ""

    removed, catalog, colours_out = [], [], []
    for it in items:
        h = it["handle"]
        cat = gc.category_of(h, it["ptype"], it["name"])
        ok, total = stock(it)
        if h in gc.HIDDEN or cat in gc.HIDDEN_CATS:
            why = gc.HIDDEN.get(h) or f"whole category taken off the site ({gc.CATS[cat][2]})"
            log.setdefault(h, today)
            removed.append((it, why, log[h]))
            status, where = "Removed", why
        elif h in combined:
            status, where = combined[h]
        else:
            status, where = "On site", f"{SITE}gioberti-{h}.html"
        catalog.append((it, status, where, ok, total))
    log = {h: d for h, d in log.items() if any(r[0]["handle"] == h for r in removed)}      # brought back -> dropped from the log

    for cfg in gc.CURATED:
        wanted = cfg.get("colours")
        if not wanted:
            continue
        for h in cfg["sources"]:
            it = next((i for i in items if i["handle"] == h), None)
            if not it:
                continue
            for c in it["all_colours"]:
                if c in wanted or gg._fabric(c) in wanted or (cfg.get("solids") and gg.is_solid(c)):
                    continue
                ok, n = stock(it, c)
                colours_out.append((cfg["name"], it, c, ok, n))

    if "--sheets-json" in sys.argv:                                    # rows for the Google Sheets copy: photos as =IMAGE() of Gioberti's own photo URL
        def src(it, colour=None):
            im = next((im for im in it["images"] if colour in (im.get("colours") or []) and not im["chart"]), None) if colour else None
            im = im or next((im for im in it["images"] if not im["chart"]), None)
            return (im["src"] + ("&" if "?" in im["src"] else "?") + "width=200") if im else ""
        pic = lambda u: f'=IMAGE("{u}")' if u else ""
        link = lambda u: f'=HYPERLINK("{u}","Open on gioberti.com")'
        data = {"Removed Products": [], "Removed Colours": [], "Full Catalog": []}
        for it, why, when in sorted(removed, key=lambda x: (cat_label(x[0]["handle"], x[0]), x[0]["name"])):
            ok, total = stock(it)
            data["Removed Products"].append([pic(src(it, it["colours"][0] if it["colours"] else None)), it["name"], it["code"], cat_label(it["handle"], it), why, when,
                                             "Sold out" if it["soldout"] else f"{ok} of {total} options in stock", ", ".join(it["all_colours"]),
                                             ", ".join(it["options"][1]["values"] if len(it["options"]) > 1 else []), it["price"], link(it["supplier_url"])])
        for site_name, it, c, ok, n in sorted(colours_out, key=lambda x: (x[0], x[1]["code"], x[2])):
            data["Removed Colours"].append([pic(src(it, c)), c, site_name, it["name"], it["code"], f"{ok} of {n}", link(it["supplier_url"])])
        for it, status, where, ok, total in sorted(catalog, key=lambda x: (x[1] != "On site", cat_label(x[0]["handle"], x[0]), x[0]["name"])):
            data["Full Catalog"].append([it["name"], it["code"], status, where, cat_label(it["handle"], it), it["ptype"], it["price"], len(it["all_colours"]),
                                         "Sold out" if it["soldout"] else f"{ok} of {total}", link(it["supplier_url"])])
        path = sys.argv[sys.argv.index("--sheets-json") + 1]
        json.dump(data, open(path, "w", encoding="utf-8"), ensure_ascii=False)
        print("wrote", path, {k: len(v) for k, v in data.items()})

    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True                               # Summary counts are formulas: Excel works them out when the file opens

    def sheet(title, headers, widths):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for i, (h, w) in enumerate(zip(headers, widths), 1):
            c = ws.cell(row=1, column=i)
            c.font = Font(name=FONT, bold=True, color="FFFFFF")
            c.fill = HEAD_FILL
            c.alignment = Alignment(vertical="center", wrap_text=True)
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.row_dimensions[1].height = 30
        ws.freeze_panes = "B2" if headers[0] == "Photo" else "A2"
        return ws

    def style_row(ws, r, link_cols=(), tall=False):
        for c in ws[r]:
            c.font = Font(name=FONT, size=10, color="0563C1" if c.column in link_cols and c.value else "000000", underline="single" if c.column in link_cols and c.value else None)
            c.alignment = Alignment(vertical="center", wrap_text=True)
            c.border = THIN
            if r % 2 == 0:
                c.fill = BAND_FILL
        if tall:
            ws.row_dimensions[r].height = 58

    # Removed Products
    ws = sheet("Removed Products", ["Photo", "Product", "Style", "Former site category", "Why removed", "Removed on", "Stock at Gioberti",
                                    "Colours", "Sizes", "Price ($)", "Gioberti page", "Photo folder"], [11, 30, 11, 22, 34, 12, 16, 46, 26, 10, 18, 18])
    for it, why, when in sorted(removed, key=lambda x: (cat_label(x[0]["handle"], x[0]), x[0]["name"])):
        ok, total = stock(it)
        sizes = it["options"][1]["values"] if len(it["options"]) > 1 else []
        ws.append([None, it["name"], it["code"], cat_label(it["handle"], it), why, when,
                   "Sold out" if it["soldout"] else f"{ok} of {total} options in stock", ", ".join(it["all_colours"]), ", ".join(sizes), it["price"],
                   "Open on gioberti.com", "Open photo folder"])
        r = ws.max_row
        ws.cell(row=r, column=11).hyperlink = it["supplier_url"]
        fo = folder(it)
        if fo:
            ws.cell(row=r, column=12).hyperlink = "file:///" + fo.replace("\\", "/")
        ws.cell(row=r, column=10).number_format = '"$"#,##0.00'
        style_row(ws, r, link_cols=(11, 12), tall=True)
        pic = thumb(colour_photo(it, it["colours"][0] if it["colours"] else None))
        if pic:
            ws.add_image(pic, f"A{r}")
    ws.auto_filter.ref = ws.dimensions

    # Removed Colours
    ws = sheet("Removed Colours", ["Photo", "Colour", "On-site listing", "Gioberti listing", "Style", "Sizes in stock", "Gioberti page", "Photo folder"],
               [11, 34, 34, 34, 11, 14, 18, 18])
    for site_name, it, c, ok, n in sorted(colours_out, key=lambda x: (x[0], x[1]["code"], x[2])):
        ws.append([None, c, site_name, it["name"], it["code"], f"{ok} of {n}", "Open on gioberti.com", "Open photo folder"])
        r = ws.max_row
        ws.cell(row=r, column=7).hyperlink = it["supplier_url"]
        fo = folder(it, c)
        if fo:
            ws.cell(row=r, column=8).hyperlink = "file:///" + fo.replace("\\", "/")
        style_row(ws, r, link_cols=(7, 8), tall=True)
        pic = thumb(colour_photo(it, c))
        if pic:
            ws.add_image(pic, f"A{r}")
    ws.auto_filter.ref = ws.dimensions

    # Full Catalog
    ws = sheet("Full Catalog", ["Product", "Style", "Status", "On the site as / why removed", "Site category", "Gioberti type", "Price ($)",
                                "Colours", "Stock at Gioberti", "Gioberti page"], [32, 11, 11, 52, 24, 22, 10, 9, 18, 18])
    for it, status, where, ok, total in sorted(catalog, key=lambda x: (x[1] != "On site", cat_label(x[0]["handle"], x[0]), x[0]["name"])):
        ws.append([it["name"], it["code"], status, where, cat_label(it["handle"], it), it["ptype"], it["price"], len(it["all_colours"]),
                   "Sold out" if it["soldout"] else f"{ok} of {total}", "Open on gioberti.com"])
        r = ws.max_row
        ws.cell(row=r, column=10).hyperlink = it["supplier_url"]
        ws.cell(row=r, column=7).number_format = '"$"#,##0.00'
        if where.startswith("http") or "(https" in where:
            ws.cell(row=r, column=4).hyperlink = re.search(r"https?://\S+?\.html", where).group(0)
        style_row(ws, r, link_cols=(4, 10) if ws.cell(row=r, column=4).hyperlink else (10,))
    ws.auto_filter.ref = ws.dimensions

    # Summary (first tab)
    ws = wb["Sheet"]
    ws.title = "Summary"
    wb.move_sheet("Summary", offset=-3)
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 70
    ws["A1"] = "Checkmatela product database"
    ws["A1"].font = Font(name=FONT, size=16, bold=True, color="1E4D3B")
    ws["A2"] = f"Built {today} from Gioberti's live catalog and the site's rules (tools/gb_catalog.py). Not published on the site."
    ws["A2"].font = Font(name=FONT, size=10, italic=True, color="555555")
    rows = [
        ("Gioberti listings (all)", "=COUNTA('Full Catalog'!A:A)-1", "Every product Gioberti sells (gift card and hidden helper items excluded)"),
        ("On the site as their own page", "=COUNTIF('Full Catalog'!C:C,\"On site\")", ""),
        ("Combined into another page", "=COUNTIF('Full Catalog'!C:C,\"Combined\")", "e.g. two Gioberti listings of the same shirt shown as one page"),
        ("Removed from the site", "=COUNTIF('Full Catalog'!C:C,\"Removed\")", "Listed on the Removed Products tab"),
        ("Colourways trimmed from listings still on the site", "=COUNTA('Removed Colours'!B:B)-1", "Listed on the Removed Colours tab"),
    ]
    for i, (label, formula, note) in enumerate(rows, start=4):
        ws.cell(row=i, column=1, value=label).font = Font(name=FONT, size=11)
        c = ws.cell(row=i, column=2, value=formula)
        c.font = Font(name=FONT, size=11, bold=True)
        ws.cell(row=i, column=3, value=note).font = Font(name=FONT, size=10, color="555555")
    ws["A11"] = "How to bring something back"
    ws["A11"].font = Font(name=FONT, size=12, bold=True, color="1E4D3B")
    tips = ["A removed product: delete its line from HIDDEN (or its category from HIDDEN_CATS) in tools/gb_catalog.py, then run tools/gen_gioberti.py - or just ask Claude.",
            "A removed colour: add it to that listing's colours in CURATED in tools/gb_catalog.py (the Gioberti listing and colour name are on the Removed Colours tab).",
            "Photos of every product, removed ones included, stay in Downloads\\Gioberti (the Photo folder links open them).",
            "To refresh this workbook with today's stock: python tools/gb_database.py (run from the checkmatela-repo folder)."]
    for i, t in enumerate(tips, start=12):
        ws.cell(row=i, column=1, value="•  " + t).font = Font(name=FONT, size=10)
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=3)
        ws.cell(row=i, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[i].height = 30

    wb.save(out)
    json.dump(log, open(log_path, "w", encoding="utf-8"), indent=1)
    print("wrote", out)
    print(f"  {len(catalog)} listings: {sum(1 for c in catalog if c[1] == 'On site')} on site, {sum(1 for c in catalog if c[1] == 'Combined')} combined, "
          f"{len(removed)} removed; {len(colours_out)} trimmed colourways")


if __name__ == "__main__":
    main()
